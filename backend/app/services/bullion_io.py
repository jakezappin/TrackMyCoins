"""Bullion spreadsheets: export a collection (or everything) to Excel, and import a CSV / Excel file of items.

The export's columns import back as they are; other spreadsheets work with headings like Description, Qty,
Weight, Unit, Purity (".999", "999", "90%", "22k"), Face value, Price paid, Date purchased, Dealer, Notes.
"""

import io
import re
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..models import BullionItem
from . import bullion as svc
from .importer import _date, _key, _money, _text, read_rows

COLUMNS = ["Inventory #", "Collection", "Description", "Packaging", "Quantity", "Weight", "Unit", "Purity",
           "Face value", "Pure ozt", "Price paid", "Date purchased", "Purchased from", "Melt value", "Sale price",
           "Sale date", "Sold to", "Notes"]

ALIASES = {
    "description": "description", "item": "description", "product": "description", "name": "description",
    "packaging": "packaging", "type": "packaging", "form": "packaging",
    "quantity": "quantity", "qty": "quantity", "pieces": "quantity", "count": "quantity",
    "weight": "weight", "weighteach": "weight", "weightperpiece": "weight", "unitweight": "weight", "size": "weight",
    "unit": "weight_unit", "units": "weight_unit", "weightunit": "weight_unit",
    "purity": "purity", "fineness": "purity", "fine": "purity", "karat": "purity",
    "facevalue": "face_value", "face": "face_value",
    "pricepaid": "price_paid", "paid": "price_paid", "cost": "price_paid", "totalcost": "price_paid",
    "purchaseprice": "price_paid",
    "datepurchased": "purchase_date", "purchasedate": "purchase_date", "purchased": "purchase_date", "date": "purchase_date",
    "purchasedfrom": "purchased_from", "dealer": "purchased_from", "seller": "purchased_from", "from": "purchased_from",
    "saleprice": "sale_price", "soldfor": "sale_price",
    "saledate": "sale_date", "datesold": "sale_date",
    "soldto": "sold_to", "buyer": "sold_to",
    "notes": "notes", "note": "notes", "comments": "notes",
}

UNIT_WORDS = {"ozt": "ozt", "oz": "ozt", "troyoz": "ozt", "troyounce": "ozt", "troyounces": "ozt", "ounce": "ozt",
              "ounces": "ozt", "g": "g", "gram": "g", "grams": "g", "gr": "g", "kg": "kg", "kilo": "kg", "kilos": "kg",
              "kilogram": "kg", "kilograms": "kg"}
KARATS = {"24k": "0.999", "22k": "0.9167", "18k": "0.75", "14k": "0.585"}


def parse_purity(v) -> Decimal | None:
    """'.999', '999', '0.999', '99.9%', '90%', '22k' -> a share between 0 and 1."""
    text = (_text(v) or "").lower().replace(" ", "").replace("fine", "")
    if not text:
        return None
    if text in KARATS:
        return Decimal(KARATS[text])
    try:
        if text.endswith("%"):
            n = Decimal(text[:-1]) / 100
        else:
            n = Decimal(text)
            if n > 1:
                n = n / (Decimal(1000) if n >= 100 else Decimal(100))  # 999 / 9999 or 90
                if n > 1:
                    n = n / 10  # 9999 -> .9999
    except InvalidOperation:
        return None
    return n if 0 < n <= 1 else None


def parse_weight(v, unit_v) -> tuple[Decimal | None, str]:
    """'1', '10 oz', '1 kilo', '31.1 g' (a unit in the weight wins over the Unit column)."""
    text = (_text(v) or "").lower()
    unit = UNIT_WORDS.get(_key(unit_v) or "", "ozt")
    m = re.match(r"^\s*([0-9]*\.?[0-9]+)\s*([a-z ]*)$", text)
    if not m:
        return None, unit
    if m.group(2).strip():
        unit = UNIT_WORDS.get(re.sub(r"[^a-z]", "", m.group(2)), unit)
    return Decimal(m.group(1)), unit


def row_to_item(raw: dict) -> tuple[dict, list[str]]:
    vals: dict = {}
    for header, value in raw.items():
        field = ALIASES.get(_key(header))
        if field and field not in vals:
            vals[field] = value
    errors: list[str] = []
    f: dict = {"description": _text(vals.get("description"))}
    if not f["description"]:
        errors.append("No description")
    purity = parse_purity(vals.get("purity"))
    if purity is None:
        errors.append(f"Couldn't read the purity: {vals.get('purity')!r} (use .999, 90%, 22k...)")
    f["purity"] = str(purity) if purity is not None else None
    face = _money(vals.get("face_value"), errors, "Face value")
    weight, unit = parse_weight(vals.get("weight"), vals.get("weight_unit"))
    if face is not None:
        f.update(face_value=str(face), weight=None, quantity=1)
    elif weight is None:
        errors.append("No weight (or face value for junk silver)")
    else:
        f.update(weight=str(weight), face_value=None)
    f["weight_unit"] = unit
    if face is None:
        try:
            f["quantity"] = int(Decimal(str(_text(vals.get("quantity")) or "1")))
        except InvalidOperation:
            errors.append(f"Quantity isn't a number: {vals.get('quantity')}")
            f["quantity"] = 1
    pack = (_text(vals.get("packaging")) or "").lower()
    f["packaging"] = next((p for p in svc.PACKAGING if p in pack), "tube" if f.get("quantity", 1) > 1 and "roll" not in pack else "single")
    if "roll" in pack:
        f["packaging"] = "roll"
    for k, label in (("price_paid", "Price paid"), ("sale_price", "Sale price")):
        m_ = _money(vals.get(k), errors, label)
        f[k] = str(m_) if m_ is not None else None
    for k, label in (("purchase_date", "Date purchased"), ("sale_date", "Sale date")):
        d = _date(vals.get(k), errors, label)
        f[k] = d.isoformat() if d else None
    for k in ("purchased_from", "sold_to", "notes"):
        f[k] = _text(vals.get(k))
    return f, errors


def read(filename: str, data: bytes) -> list[tuple[str, int, dict]]:
    return read_rows(filename, data, sheets=("Bullion", "Sold"))


def export_xlsx(items: list[BullionItem], spots: dict[str, Decimal | None]) -> bytes:
    wb = Workbook()
    held = wb.active
    held.title = "Bullion"
    sold = wb.create_sheet("Sold")
    for ws in (held, sold):
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for i in sorted(items, key=lambda i: i.inventory_number):
        oz = svc.pure_oz(i)
        spot = spots.get(i.collection.metal)
        (sold if svc.is_sold(i) else held).append([
            i.inventory_number, i.collection.name, i.description, i.packaging, i.quantity,
            float(i.weight) if i.weight is not None else None, i.weight_unit, float(i.purity),
            float(i.face_value) if i.face_value is not None else None, float(oz),
            float(i.price_paid) if i.price_paid is not None else None, i.purchase_date, i.purchased_from,
            float((oz * spot).quantize(svc.CENTS)) if spot is not None else None,
            float(i.sale_price) if i.sale_price is not None else None, i.sale_date, i.sold_to, i.notes,
        ])
    for ws in (held, sold):
        for col, width in zip("ABCDEFGHIJKLMNOPQR", (12, 18, 28, 10, 9, 9, 6, 8, 10, 10, 11, 13, 18, 11, 11, 11, 16, 30)):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for idx in (10, 13, 14):  # money
                row[idx].number_format = '"$"#,##0.00'
            for idx in (11, 15):
                row[idx].number_format = "yyyy-mm-dd"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


TEMPLATE_CSV = (
    "Description,Packaging,Quantity,Weight,Unit,Purity,Face value,Price paid,Date purchased,Purchased from,Notes\n"
    "Buffalo Round,Tube,20,1,ozt,.999,,1150.00,2025-03-15,APMEX,\n"
    "Silvertown Bar,Single,1,10,ozt,.999,,330.00,2025-04-01,Local coin shop,\n"
    "90% dimes,,,,,90%,10.00,300.00,2025-05-10,Show,\n"
)

