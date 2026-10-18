"""Spreadsheet cleanup: take a messy collection spreadsheet and return it in TrackMyCoins' import format.

1. Columns are matched to fields by their headings (the same names the importer knows); the user can re-map any
   column, ignore it, or fold it into Notes.
2. A run-together description ("1881-S Morgan Dollar PCGS MS65+ CAC 12345678") is split into date, service, grade,
   CAC and cert when the file has no separate columns for them.
3. Every row goes through the importer's own reading (so cleaning and importing agree), and is written back in
   standard form: "1881-S", "MS65+ DMPL", "PCGS", "$1,250" -> "1250.00", "3/15/25" -> "2025-03-15"...
4. What changed, and what still needs fixing (unknown series, impossible grades, duplicate certs), is listed per row.
"""

import io
import re
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import CatalogEntry, CoinType
from ..series import BY_SLUG
from . import details, series_match
from .coins import _cert_key, find_duplicate_cert
from .import_columns import COLUMNS, HEADERS
from .importer import ALIASES, _key, _text, no_formulas, row_to_fields

NOTES = "notes_append"  # mapping choice: keep this column's text in Notes
# Fields the importer reads that the standard columns don't cover, with the heading they get.
EXTRA = {"year": "Year", "mint_mark": "Mint", "strike_designation": "Designation", "cac": "CAC", "ogp": "OGP",
         "details": "Details", "details_problem": "Problem", "details_pct": "Details %", "toner_premium": "Toner premium",
         "holder_premium": "Holder premium", "value_source": "Value source", "estimated_value": "Estimated value",
         "inventory_hint": "Inventory #", "purchase_value_from": "Value at purchase from"}
FIELD_HEADER = {**{f: h for h, f, *_ in COLUMNS}, **EXTRA}
CHOICES = [{"field": f, "label": h} for h, f, *_ in COLUMNS] + [{"field": f, "label": h} for f, h in EXTRA.items()] + [
    {"field": NOTES, "label": "Add to Notes"}, {"field": "", "label": "Ignore"}]

_SERVICE = re.compile(r"\b(PCGS|NGC|ANACS|CACG|ICG)\b", re.I)
_GRADE = re.compile(r"\b(MS|PR|PF|SP|AU|XF|EF|VF|VG|AG|FR|PO|F|G)[\s-]?(\d{1,2})(\s?\+)?(\s?(?:★|\*))?(?![\d/])", re.I)
_CERT = re.compile(r"(?:#\s*)?\b(\d{7}-\d{3}|\d{6,10})\b")
_CAC = re.compile(r"\bCAC\b(?:\s+(green|gold))?|\b(green|gold)\s+(?:cac|bean|sticker)\b", re.I)
_DESIG = re.compile(r"\b(DMPL|DPL|PL|DCAM|UCAM|CAM|RD|RB|BN|FBL|FB|FS|FH|FT)\b")


def auto_field(header: str) -> str | None:
    return ALIASES.get(_key(header))


def split_description(text: str) -> tuple[str, dict]:
    """("1881-S Morgan Dollar", {"grading_service": "PCGS", "grade": "MS65+ DMPL", "cac_sticker": "Green",
    "cert_number": "12345678"}) from "1881-S Morgan Dollar PCGS MS65+ DMPL CAC Green #12345678"."""
    found: dict = {}
    rest = text
    if (m := _SERVICE.search(rest)):
        found["grading_service"] = m.group(1).upper()
        rest = rest[: m.start()] + " " + rest[m.end():]
    if (m := _CAC.search(rest)):
        found["cac_sticker"] = (m.group(1) or m.group(2) or "yes").title()
        rest = rest[: m.start()] + " " + rest[m.end():]
    # The cert (6+ digits, or NGC's 1234567-001) first, so it isn't taken for part of a Details problem.
    if (m := _CERT.search(rest)):
        found["cert_number"] = m.group(1)
        rest = rest[: m.start()] + " " + rest[m.end():]
    is_details, _ = details.detect(rest)
    if (m := _GRADE.search(rest)):
        grade = f"{m.group(1).upper().replace('PF', 'PR').replace('EF', 'XF')}{m.group(2)}"
        grade += "+" if m.group(3) else ""
        grade += "★" if m.group(4) else ""
        tail = rest[m.end():]
        desig = _DESIG.findall(tail[:20])
        if desig:
            grade += " " + " ".join(desig)
            tail = _DESIG.sub(" ", tail, count=len(desig))
        if is_details:
            found_d = re.search(r"details?.*$", tail, re.I)
            grade += " Details" + ((" " + re.sub(r"^details?\W*", "", found_d.group(0), flags=re.I)) if found_d else "")
            tail = tail[: found_d.start()] if found_d else tail
        found["grade"] = grade.strip()
        rest = rest[: m.start()] + " " + tail
    return " ".join(rest.replace(" ,", ",").split()).strip(" ,-·"), found


def _money(v) -> str:
    return f"{Decimal(str(v)):.2f}" if v not in (None, "") else ""


def _grade_text(f: dict) -> str:
    if f.get("details"):
        return details.label(_NS(f)).replace(" · ", " ")
    g = f"{f.get('grade_prefix') or ''}{f.get('grade') or ''}"
    if g and f.get("grade_plus"):
        g += "+"
    if g and f.get("grade_star"):
        g += "★"
    if f.get("strike_designation"):
        g = f"{g} {f['strike_designation']}".strip()
    return g


class _NS:
    def __init__(self, d: dict):
        self.__dict__.update(d)


def _join(pairs) -> str:
    return ", ".join(label for on, label in pairs if on)


def _series_name(db: Session, f: dict, hints: dict, by_pcgs: dict[str, int], names: dict[int, str]) -> tuple[str | None, str | None]:
    """(series name, why not) for a row, the same ways the importer works it out."""
    s = series_match.from_inventory_number(hints.get("inventory"))
    if s:
        return s.name, None
    type_id = by_pcgs.get(str(f.get("pcgs_number") or ""))
    if type_id:
        return names.get(type_id), None
    m = series_match.from_text(hints.get("text"), f.get("year"))
    if m.series:
        return m.series.name, None
    return None, m.reason or ("No series: add a Type column (e.g. Morgan Dollar) or put it after the date"
                              if not hints.get("text") else f"Didn't recognize {hints.get('text')!r} as a series")


def clean(db: Session, user_id: int, rows: list[tuple[str, int, dict]], mapping: dict[str, str] | None = None) -> dict:
    mapping = mapping or {}
    headers: list[str] = []
    for _, _, raw in rows:
        for h in raw:
            if h not in headers:
                headers.append(h)
    fields_for = {h: (mapping[h] if h in mapping else auto_field(h)) for h in headers}
    samples = {h: [str(raw[h]) for _, _, raw in rows if _text(raw.get(h))][:3] for h in headers}

    pcgs_numbers = set()
    names = {t.id: (BY_SLUG[t.slug].name if t.slug in BY_SLUG else t.name) for t in db.scalars(select(CoinType))}
    out_rows, used = [], set()
    first_pass = []
    for sheet, num, raw in rows:
        canon: dict[str, object] = {}
        originals: dict[str, list[str]] = {}
        notes_extra = []
        for h in headers:
            v = raw.get(h)
            field = fields_for.get(h)
            text = _text(v)
            if field == NOTES:
                if text:
                    notes_extra.append(f"{h}: {text}")
                continue
            if not field or text is None:
                continue
            head = FIELD_HEADER.get(field, h)
            originals.setdefault(head, []).append(text)
            if head not in canon:
                canon[head] = v
        if notes_extra:
            canon["Notes"] = "; ".join(x for x in [_text(canon.get("Notes")), *notes_extra] if x)
        # A run-together description: split out what has no column of its own.
        split_from = None
        date_text = _text(canon.get("Date"))
        if date_text:
            remaining, found = split_description(date_text)
            took = {k: v for k, v in found.items() if not _text(canon.get(FIELD_HEADER[k]))}
            if took:
                canon["Date"] = remaining
                for k, v in took.items():
                    canon[FIELD_HEADER[k]] = v
                split_from = date_text
        f, valuation, errors, hints = row_to_fields(canon, sold_sheet=sheet.lower() == "sold")
        if f.get("pcgs_number"):
            pcgs_numbers.add(str(f["pcgs_number"]))
        first_pass.append((sheet, num, canon, originals, split_from, f, valuation, errors, hints))

    by_pcgs: dict[str, int] = {}
    if pcgs_numbers:
        for n, t in db.execute(select(CatalogEntry.pcgs_number, CatalogEntry.coin_type_id)
                               .where(CatalogEntry.pcgs_number.in_(pcgs_numbers))):
            by_pcgs.setdefault(n, t)

    seen_certs: dict[tuple, int] = {}
    for sheet, num, canon, originals, split_from, f, valuation, errors, hints in first_pass:
        problems = list(errors)
        series, why = _series_name(db, f, hints, by_pcgs, names)
        if series is None and why:
            problems.append(why)
        mint = f.get("mint_mark") or "P"
        cleaned = {
            "Type": series or _text(canon.get("Type")) or "",
            "Date": f"{f['year']}" + ("" if mint == "P" else f"-{mint}") if f.get("year") else _text(canon.get("Date")) or "",
            "Variety": f.get("variety") or "",
            "VAM": f.get("vam") or "",
            "Grading service": "Raw" if f.get("grading_service") == "RAW" else (f.get("grading_service") or ""),
            "Cert #": f.get("cert_number") or "",
            "Grade": _grade_text(f),
            "CAC sticker": (f.get("cac_sticker") or "").title() or ("Yes" if f.get("cac") else ""),
            "Collector slab": _join([(f.get("slab_soapbox"), "Soapbox"), (f.get("slab_ngc_fatty"), "NGC Fatty"),
                                     (f.get("slab_rattler"), "Rattler"), (f.get("slab_ogh"), "OGH")]),
            "Packaging": _join([(f.get("gsa_hard_pack"), "GSA Hard Pack"), (f.get("gsa_soft_pack"), "GSA Soft Pack"),
                                (f.get("ike_blue_pack"), "Blue Pack"), (f.get("ike_brown_pack"), "Brown Pack"),
                                (f.get("ogp"), "OGP"), (f.get("capital_holder"), "Capital Holder")]),
            "Toning": _join([(f.get("toning_obverse"), "Obverse"), (f.get("toning_reverse"), "Reverse"),
                             (f.get("toning_dual"), "Dual"), (f.get("toning_rainbow"), "Rainbow"), (f.get("toning_rim"), "Rim")]),
            "Error": f.get("error_description") or "",
            "PCGS #": f.get("pcgs_number") or "",
            "Price paid": _money(f.get("price_paid")),
            "Date purchased": f.get("purchase_date") or "",
            "Purchased from": f.get("purchased_from") or "",
            "Est. value at purchase": _money(f.get("purchase_value_custom")),
            "Sale price": _money(f.get("sale_price")),
            "Sale date": f.get("sale_date") or "",
            "Sold to": f.get("sold_to") or "",
            "Notes": f.get("notes") or "",
            "Stored in": hints.get("storage") or "",
            "Position": hints.get("position") or "",
            "Status": {"grading": "At Grading", "inventory": "In Inventory"}.get(hints.get("status") or "", ""),
            "Submission": hints.get("submission") or "",
            "Submitted to": hints.get("submit_company") or "",
            "Toner premium": _money(valuation.get("toner_premium")),
            "Holder premium": _money(valuation.get("holder_premium")),
            "Value source": "Own value" if valuation.get("value_source") == "custom" else "",
            "Estimated value": _money(valuation.get("custom_value")),
        }
        # Certs twice in the file, or already in the collection.
        if f.get("cert_number") and f.get("grading_service") not in (None, "RAW") and not (f.get("sale_price") or f.get("sale_date")):
            key = (f["grading_service"], _cert_key(f["cert_number"]))
            if key in seen_certs:
                problems.append(f"Same cert as row {seen_certs[key]}")
            else:
                seen_certs[key] = num
            dup = find_duplicate_cert(db, user_id, f["grading_service"], f["cert_number"])
            if dup:
                problems.append(f"Already in your collection as {dup.inventory_number}")
        changes = []
        for head, after in cleaned.items():
            before = "; ".join(originals.get(head, []))
            if head == "Date" and split_from:
                before = split_from
            if _norm(before) != _norm(after) and (before or after):
                changes.append({"column": head, "before": before, "after": after})
        used |= {h for h, v in cleaned.items() if v}
        notes = []
        if hints.get("service_guessed"):
            notes.append(f"Grading service {cleaned['Grading service']} worked out from the cert number's format. Check it.")
        out_rows.append({"sheet": sheet, "row": num, "cleaned": cleaned, "changes": changes, "problems": problems,
                         "notes": notes})

    columns = [h for h in HEADERS if h in used or h in ("Type", "Date", "Grading service", "Cert #", "Grade")]
    columns += [h for h in ("Toner premium", "Holder premium", "Value source", "Estimated value") if h in used]
    return {
        "headers": [{"header": h, "field": fields_for.get(h) or "", "auto": auto_field(h) or "", "samples": samples[h]}
                    for h in headers],
        "choices": CHOICES,
        "columns": columns,
        "rows": out_rows,
        "counts": {"rows": len(out_rows), "changed": sum(1 for r in out_rows if r["changes"]),
                   "problems": sum(1 for r in out_rows if r["problems"])},
    }


def _norm(v: str) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def to_xlsx(result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Coins"
    cols = result["columns"]
    ws.append(cols)
    for r in result["rows"]:
        ws.append([r["cleaned"].get(c, "") for c in cols])
    head = PatternFill("solid", fgColor="1F3A5F")
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(30, len(c) + 8))
    ws.freeze_panes = "A2"

    ch = wb.create_sheet("Changes")
    ch.append(["Row in your file", "Column", "Was", "Now"])
    for r in result["rows"]:
        for c in r["changes"]:
            ch.append([r["row"], c["column"], c["before"], c["after"]])
    pr = wb.create_sheet("Problems")
    pr.append(["Row in your file", "Problem"])
    for r in result["rows"]:
        for p in r["problems"]:
            pr.append([r["row"], p])
    for sheet in (ch, pr):
        for cell in sheet[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
        for col, width in zip("ABCD", (16, 22, 40, 40)):
            sheet.column_dimensions[col].width = width
    no_formulas(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

