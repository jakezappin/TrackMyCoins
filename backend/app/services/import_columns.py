"""The import spreadsheet's columns: one list for the template, its Instructions sheet, the Import page's column
guide and the cleanup tool's output. Headers here are read back by the importer (services/importer.ALIASES)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REQUIRED, SOMETIMES, OPTIONAL = "Required", "Sometimes required", "Optional"

# (header, field, need, what to enter, examples)
COLUMNS: list[tuple[str, str, str, str, str]] = [
    ("Type", "series_hint", SOMETIMES,
     "The series. Needed when sorting into collections automatically, unless the Date says it (1881-S Morgan "
     "Dollar), or there's a PCGS # or TrackMyCoins inventory #.", "Morgan Dollar; Roosevelt Dime; American Silver Eagle"),
    ("Date", "date", REQUIRED, "Year and mint mark (no mint mark = Philadelphia). Overdates and over-mintmarks are "
     "fine. Or use separate Year and Mint columns.", "1881-S; 1921; 1880/79-CC; 1882-O/S"),
    ("Variety", "variety", OPTIONAL, "Variety, issue or pedigree.", "8TF; 7TF Reverse of 1878; Type 2"),
    ("VAM", "vam", OPTIONAL, "VAM number (Morgan and Peace dollars).", "VAM-3"),
    ("Grading service", "grading_service", SOMETIMES, "Needed when there's a Cert #. Leave blank (or Raw) for "
     "ungraded coins.", "PCGS; NGC; ANACS; CACG; ICG; Raw"),
    ("Cert #", "cert_number", SOMETIMES, "The slab's certification number. Needed for graded coins.",
     "12345678; 1234567-001"),
    ("Grade", "grade", OPTIONAL, "Grade with any plus, star or designation. Details grades: add Details and the "
     "problem.", "MS65+; PR69 DCAM; MS64 DMPL; AU55 Details Cleaned; Raw XF"),
    ("CAC sticker", "cac_sticker", OPTIONAL, "Green or Gold (or Yes).", "Green; Gold"),
    ("Collector slab", "slab", OPTIONAL, "Special holders.", "Rattler; NGC Fatty; OGH; Soapbox"),
    ("Packaging", "packaging", OPTIONAL, "Original packaging.", "OGP; GSA Hard Pack; Blue Pack; Capital Holder"),
    ("Toning", "toning", OPTIONAL, "Which sides, or Yes.", "Obverse; Reverse, Rainbow; Rim"),
    ("Error", "error_description", OPTIONAL, "For mint errors: describe it (the coin is then valued by you).",
     "Off-center 20%; Double struck"),
    ("PCGS #", "pcgs_number", OPTIONAL, "PCGS coin number, for pricing. Found automatically when left blank.",
     "7130"),
    ("Price paid", "price_paid", OPTIONAL, "What you paid.", "250; $1,250.00"),
    ("Date purchased", "purchase_date", OPTIONAL, "When you bought it.", "2025-03-15; 3/15/2025"),
    ("Purchased from", "purchased_from", OPTIONAL, "Dealer, show or auction.", "Heritage; Local coin shop"),
    ("Est. value at purchase", "purchase_value", OPTIONAL, "Your own figure for what it was worth when bought. "
     "Leave blank to have it worked out.", "300"),
    ("Sale price", "sale_price", OPTIONAL, "If sold. A sale price or date files it as sold.", "400"),
    ("Sale date", "sale_date", OPTIONAL, "If sold.", "2026-01-05"),
    ("Sold to", "sold_to", OPTIONAL, "If sold.", "eBay buyer"),
    ("Notes", "notes", OPTIONAL, "Anything else.", ""),
    ("Stored in", "storage_name", OPTIONAL, "Where it's kept. A storage spot by this name is created if you don't "
     "have one.", "Dansco Morgan Album; Box #12"),
    ("Position", "storage_position", OPTIONAL, "Where in it (needs Stored in).", "Page 2, hole 9; Row 3"),
    ("Status", "status", OPTIONAL, "In Inventory (the default) or At Grading.", "In Inventory; At Grading"),
    ("Submission", "submission", OPTIONAL, "At Grading: the Grading Tracker submission's # or name.", "#3; Morgans for NGC"),
    ("Submitted to", "submit_company", OPTIONAL, "At Grading without a Submission: the grader (a new submission "
     "is started for it).", "NGC; PCGS; CAC"),
]
HEADERS = [c[0] for c in COLUMNS]

EXAMPLES = [
    {"Type": "Morgan Dollar", "Date": "1881-S", "Grading service": "PCGS", "Cert #": "12345678",
     "Grade": "MS65+ DMPL", "CAC sticker": "Green", "Price paid": "250.00", "Date purchased": "2025-03-15",
     "Purchased from": "Heritage", "Stored in": "Box #12", "Position": "Row 1"},
    {"Type": "Morgan Dollar", "Date": "1878", "Variety": "7TF Reverse of 1878", "VAM": "VAM-131",
     "Grading service": "NGC", "Cert #": "1234567-001", "Grade": "MS63", "Toning": "Obverse, Rainbow",
     "Price paid": "180.00", "Status": "At Grading", "Submitted to": "PCGS"},
    {"Type": "Roosevelt Dime", "Date": "1962-D", "Grading service": "Raw", "Grade": "AU55",
     "Price paid": "4.00", "Notes": "From grandpa's jar"},
    {"Type": "Peace Dollar", "Date": "1922", "Grading service": "PCGS", "Cert #": "87654321", "Grade": "MS64",
     "Error": "Off-center 10%", "Price paid": "450.00", "Sale price": "600.00", "Sale date": "2026-02-01"},
]


def guide() -> list[dict]:
    return [{"header": h, "need": need, "about": about, "examples": ex} for h, _f, need, about, ex in COLUMNS]


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Coins"
    ws.append(HEADERS)
    for ex in EXAMPLES:
        ws.append([ex.get(h, "") for h in HEADERS])
    head = PatternFill("solid", fgColor="1F3A5F")
    need_fill = {REQUIRED: PatternFill("solid", fgColor="B3261E"), SOMETIMES: PatternFill("solid", fgColor="9A6200")}
    for i, (h, _f, need, _a, _e) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = need_fill.get(need, head)
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(28, len(h) + 6))
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Instructions")
    info.append(["TrackMyCoins import template"])
    info["A1"].font = Font(bold=True, size=14)
    info.append(["One coin per row on the Coins sheet, headings in the first row. Delete the example rows. Columns "
                 "can be in any order, and ones you don't use can be left out. Red headings are required; amber "
                 "ones are required in some cases."])
    info.append([])
    info.append(["Column", "Required?", "What to enter", "Examples"])
    for cell in info[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
    for h, _f, need, about, ex in COLUMNS:
        info.append([h, need, about, ex])
    for col, width in zip("ABCD", (24, 20, 80, 46)):
        info.column_dimensions[col].width = width
    for row in info.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
