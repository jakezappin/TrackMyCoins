"""Collection status (checklist) and Excel export."""

import io
import re
from datetime import date
from typing import Literal

import httpx
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session, object_session, selectinload

from ..auth import current_user
from ..config import get_settings
from ..db import get_db
from ..models import Coin, CoinType, Collection, User
from ..schemas import CoinOut
from ..services import checklist, details, insurance, purchase_value, upgrades
from ..services.importer import no_formulas
from ..services.coins import CPG, coin_out, collection_summary, latest_quote, on_hand, sold, value_estimate
from ..services.greysheet import PricingError
from .collections import _type_label, get_owned_collection

router = APIRouter(prefix="/api/collections", tags=["reports"])


class StatusRow(BaseModel):
    label: str
    year: int
    mint_mark: str
    variety: str | None
    coins: list[CoinOut]


class StatusOut(BaseModel):
    mode: str  # "full" | "type"
    modes: list[str]  # modes this series offers
    total: int
    filled: int
    rows: list[StatusRow]
    others: list[CoinOut]
    extras: list[StatusRow] = []  # specialty sets (mint & proof sets): listed, not counted


def _build(db: Session, col: Collection, mode: str = "full") -> checklist.Checklist:
    try:
        return checklist.build(db, col, mode)
    except PricingError as e:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Couldn't load the catalog for this series: {e}")
    except httpx.HTTPError as e:
        raise HTTPException(status.HTTP_502_BAD_GATEWAY, f"Greysheet error: {e}")


@router.get("/{collection_id}/status", response_model=StatusOut)
def collection_status(
    collection_id: int,
    mode: Literal["full", "type"] = "full",
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    col = get_owned_collection(db, user, collection_id)
    modes = checklist.modes_for(col.coin_type.slug)
    if mode not in modes:
        mode = "full"
    cl = _build(db, col, mode)
    return StatusOut(
        mode=mode,
        modes=modes,
        total=len(cl.rows),
        filled=cl.filled,
        rows=[
            StatusRow(label=r.label, year=r.year, mint_mark=r.mint_mark, variety=r.variety, coins=[coin_out(c) for c in r.coins])
            for r in cl.rows
        ],
        others=[coin_out(c) for c in cl.others],
        extras=[
            StatusRow(label=r.label, year=r.year, mint_mark=r.mint_mark, variety=r.variety, coins=[coin_out(c) for c in r.coins])
            for r in cl.extras
        ],
    )


@router.get("/{collection_id}/upgrades")
def collection_upgrades(
    collection_id: int,
    mode: Literal["full", "type"] = "full",
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    """Next grade up for each coin in the set and what it would cost (saved prices only, no API calls)."""
    col = get_owned_collection(db, user, collection_id)
    try:
        return upgrades.find(db, col, mode if mode in checklist.modes_for(col.coin_type.slug) else "full")
    except PricingError as e:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Couldn't load the catalog for this series: {e}")
    except httpx.HTTPError as e:
        raise HTTPException(status.HTTP_502_BAD_GATEWAY, f"Greysheet error: {e}")


# ---- Excel export ----------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MISSING_FILL = PatternFill("solid", fgColor="FBE5E3")
MONEY = '"$"#,##0.00'


def _grade(c: Coin) -> str:
    if c.details:
        g = details.label(c)  # "AU55 Details · Cleaned" (imports back as a Details grade)
    else:
        g = f"{c.grade_prefix or ''}{c.grade or ''}{'+' if c.grade_plus else ''}{'★' if c.grade_star else ''}"
    if c.strike_designation and g:
        g = f"{g} {c.strike_designation}"
    if c.grading_service == "RAW":
        return f"Raw {g}".strip()
    return g


def _date_mint(c: Coin) -> str:
    return f"{c.year}" if c.mint_mark == "P" else f"{c.year}-{c.mint_mark}"


def _yes(v: bool) -> str:
    return "Yes" if v else ""


def _join(pairs: list[tuple[bool, str]]) -> str:
    return ", ".join(label for on, label in pairs if on)


def _purchase_estimate(c: Coin):
    return purchase_value.estimate(object_session(c), c)


COIN_COLUMNS = [
    ("Inventory #", 12, lambda c: c.inventory_number),
    ("Date", 10, _date_mint),
    ("Year", 7, lambda c: c.year),
    ("Mint", 6, lambda c: c.mint_mark),
    ("Variety", 18, lambda c: c.variety),
    ("VAM", 10, lambda c: c.vam),
    ("Grading service", 10, lambda c: c.service_name),
    ("Cert #", 14, lambda c: c.cert_number),
    ("Grade", 13, _grade),
    ("CAC", 6, lambda c: _yes(c.cac)),
    ("CAC sticker", 10, lambda c: (c.cac_sticker or "").title()),
    ("Collector slab", 14, lambda c: _join([(c.slab_soapbox, "Soapbox"), (c.slab_ngc_fatty, "NGC Fatty"),
                                            (c.slab_rattler, "Rattler"), (c.slab_ogh, "OGH")])),
    ("Packaging", 14, lambda c: _join([(c.gsa_hard_pack, "GSA Hard Pack"), (c.gsa_soft_pack, "GSA Soft Pack"),
                                       (c.ike_blue_pack, "Blue Pack"), (c.ike_brown_pack, "Brown Pack"),
                                       (c.ogp, "OGP"), (c.capital_holder, "Capital Holder")])),
    ("Toning", 14, lambda c: _join([(c.toning_obverse, "Obverse"), (c.toning_reverse, "Reverse"),
                                    (c.toning_dual, "Dual"), (c.toning_rainbow, "Rainbow"), (c.toning_rim, "Rim")])),
    ("Error", 20, lambda c: c.error_description if c.is_error else None),
    ("PCGS #", 9, lambda c: c.pcgs_number),
    ("Price paid", 12, lambda c: c.price_paid),
    ("Date purchased", 12, lambda c: c.purchase_date),
    ("Purchased from", 18, lambda c: c.purchased_from),
    ("Est. value at purchase", 14, lambda c: _purchase_estimate(c).value),
    ("Value at purchase from", 22, lambda c: _purchase_estimate(c).detail or None),
    ("CPG value", 12, lambda c: (q.value if (q := latest_quote(c, CPG)) else None)),
    ("CPG as of", 12, lambda c: (q.fetched_at.date() if (q := latest_quote(c, CPG)) else None)),
    ("Toner premium", 12, lambda c: c.toner_premium),
    ("Holder premium", 12, lambda c: c.holder_premium),
    ("Details %", 9, lambda c: details.pct(c) if c.details else None),
    ("Value source", 10, lambda c: "Own value" if c.value_source == "custom" else "CPG"),
    ("Estimated value", 14, lambda c: value_estimate(c)),
    ("Stored in", 16, lambda c: c.storage.name if c.storage else None),
    ("Position", 12, lambda c: c.storage_position),
    ("Photos", 7, lambda c: len(c.photos)),
    ("Notes", 30, lambda c: c.notes),
]
SOLD_COLUMNS = COIN_COLUMNS[:9] + [
    ("Price paid", 12, lambda c: c.price_paid),
    ("Sale price", 12, lambda c: c.sale_price),
    ("Sale date", 12, lambda c: c.sale_date),
    ("Sold to", 18, lambda c: c.sold_to),
    ("Profit / loss", 12, lambda c: (c.sale_price - c.price_paid) if c.sale_price is not None and c.price_paid is not None else None),
]


SALE_ONLY = [c for c in SOLD_COLUMNS if c[0] not in {n for n, _, _ in COIN_COLUMNS}]
ALL_COLUMN_NAMES = [n for n, _, _ in COIN_COLUMNS] + [n for n, _, _ in SALE_ONLY]


def _pick(columns, chosen: set[str] | None):
    return columns if chosen is None else [c for c in columns if c[0] in chosen]


def _sheet(ws, columns, coins):
    ws.append([name for name, _, _ in columns])
    for coin in coins:
        ws.append([fn(coin) for _, _, fn in columns])
    _style(ws, [w for _, w, _ in columns], money_cols={n for n, _, _ in columns if n in (
        "Price paid", "CPG value", "Sale price", "Profit / loss", "Toner premium", "Holder premium", "Estimated value",
        "Est. value at purchase")}, header=[n for n, _, _ in columns])


def _style(ws, widths, money_cols=(), header=()):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for i, name in enumerate(header, start=1):
        fmt = MONEY if name in money_cols else "yyyy-mm-dd" if "date" in name.lower() or name == "CPG as of" else None
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = fmt


@router.get("/{collection_id}/export")
def export_collection(
    collection_id: int,
    columns: list[str] | None = Query(None, description="Column names to include (default: all)"),
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
):
    col = get_owned_collection(db, user, collection_id)
    chosen = None
    if columns:
        chosen = {c for c in columns if c in ALL_COLUMN_NAMES}
        if not chosen:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Choose at least one column to export")
    load = (selectinload(Coin.photos), selectinload(Coin.price_quotes))
    def order(coins):
        return sorted(coins, key=lambda c: (c.year, checklist.MINT_ORDER.get(c.mint_mark, 9), c.inventory_number))

    held = order(db.scalars(on_hand(select(Coin).where(Coin.collection_id == col.id)).options(*load)).all())
    gone = order(db.scalars(sold(select(Coin).where(Coin.collection_id == col.id)).options(*load)).all())

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = collection_summary(db, col.id)
    discount = int(get_settings().wholesale_discount * 100)
    for row in [
        ("Collection", col.name),
        ("Type", _type_label(col.coin_type, col.finish)),
        ("Exported", date.today()),
        ("Coins on hand", s.coin_count),
        ("Total price paid", s.total_price_paid),
        ("Total estimated value", s.total_estimated_value),
        ("  of which CPG", s.total_cpg_value),
        (f"Est. wholesale (estimated value − {discount}%)", s.est_wholesale_value),
        ("Coins valued at your own figure", s.custom_count),
        ("Coins with a toner premium", s.toner_count),
        ("Coins with a holder premium", s.holder_count),
        ("Coins without a value yet", s.unpriced_count),
        ("Coins sold from this collection", len(gone)),
    ]:
        ws.append(row)
    for r in (5, 6, 7, 8):
        ws.cell(r, 2).number_format = MONEY
    ws.cell(3, 2).number_format = "yyyy-mm-dd"
    for r in range(1, ws.max_row + 1):
        ws.cell(r, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 32, 36

    _sheet(wb.create_sheet("Coins"), _pick(COIN_COLUMNS, chosen) or COIN_COLUMNS[:1], held)

    # Checklist(s), if the catalog can be reached; the export still works without them.
    for mode in checklist.modes_for(col.coin_type.slug):
        try:
            cl = checklist.build(db, col, mode)
        except (PricingError, httpx.HTTPError):
            break
        ws = wb.create_sheet("Checklist" if mode == "full" else "Type Set")
        ws.append(["Date" if mode == "full" else "Design", "Have", "Best grade", "Coins"])
        for r in cl.rows:
            best = r.coins[0] if r.coins else None
            ws.append([
                r.label,
                "Yes" if r.coins else "Missing",
                _grade(best) if best else "",
                ", ".join(f"{c.inventory_number} ({_grade(c)})" for c in r.coins),
            ])
            if not r.coins:
                for cell in ws[ws.max_row]:
                    cell.fill = MISSING_FILL
        ws.append([])
        ws.append([f"{cl.filled} of {len(cl.rows)} filled"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        _style(ws, [16 if mode == "full" else 34, 10, 16, 60])

    if gone:
        _sheet(wb.create_sheet("Sold"), _pick(SOLD_COLUMNS, chosen) or SOLD_COLUMNS[:1], gone)

    buf = io.BytesIO()
    no_formulas(wb)
    wb.save(buf)
    buf.seek(0)
    filename = re.sub(r"[^A-Za-z0-9 _-]+", "", col.name).strip().replace(" ", "_") or "collection"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}_{date.today():%Y-%m-%d}.xlsx"'},
    )


# ---- Insurance report (PDF) -------------------------------------------------------------------

insurance_router = APIRouter(prefix="/api/reports", tags=["reports"])


@insurance_router.get("/export-columns")
def export_columns(_: User = Depends(current_user)):
    """Columns the Excel export can include. Sale columns only appear on the Sold sheet."""
    sale = {n for n, _, _ in SALE_ONLY}
    return [{"name": n, "sold_only": n in sale} for n in ALL_COLUMN_NAMES]


@insurance_router.get("/insurance")
def insurance_report(collection_id: int | None = None, db: Session = Depends(get_db), user: User = Depends(current_user)):
    """PDF inventory of every coin on hand (or one collection), for insurance."""
    if collection_id is not None:
        cols = [get_owned_collection(db, user, collection_id)]
        scope = cols[0].name
    else:
        cols = db.scalars(select(Collection).where(Collection.user_id == user.id).order_by(Collection.name)).all()
        scope = "All collections"
    load = (selectinload(Coin.photos), selectinload(Coin.price_quotes))
    groups = []
    for col in cols:
        coins = db.scalars(on_hand(select(Coin).where(Coin.collection_id == col.id)).options(*load)).all()
        coins = sorted(coins, key=lambda c: (c.year, checklist.MINT_ORDER.get(c.mint_mark, 9), c.inventory_number))
        if coins:
            groups.append((col, coins))
    type_names = {t.id: t.name for t in db.scalars(select(CoinType)).all()}
    bullion, spots = [], {}
    if collection_id is None:  # the whole inventory: bullion too, valued at today's spot
        from ..models import BullionCollection, BullionItem
        from ..services import bullion as bullion_svc

        for bcol in db.scalars(select(BullionCollection).where(BullionCollection.user_id == user.id)
                               .order_by(BullionCollection.metal, BullionCollection.name)):
            items = [i for i in db.scalars(select(BullionItem).where(BullionItem.collection_id == bcol.id)
                                           .options(selectinload(BullionItem.photos)).order_by(BullionItem.inventory_number))
                     if not bullion_svc.is_sold(i)]
            if items:
                bullion.append((bcol, items))
                spots.setdefault(bcol.metal, bullion_svc.live_spot(bcol.metal))
    pdf = insurance.build_pdf(user, groups, type_names, scope, bullion, spots)
    filename = re.sub(r"[^A-Za-z0-9 _-]+", "", scope).strip().replace(" ", "_") or "collection"
    return StreamingResponse(
        io.BytesIO(pdf),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="Insurance_Inventory_{filename}_{date.today():%Y-%m-%d}.pdf"'},
    )
