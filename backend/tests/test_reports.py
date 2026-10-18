"""Collection status (checklist) and Excel export, against the fake Greysheet catalog."""

import io

from openpyxl import load_workbook

from test_api import coin_payload
from test_catalog import fake_greysheet, type_id  # noqa: F401  (autouse fixture)


def _collection(c, slug, finish="MS", name=None):
    r = c.post("/api/collections", json={"name": name or f"{slug} {finish}", "coin_type_id": type_id(c, slug), "finish": finish})
    assert r.status_code == 201, r.text
    return r.json()


def test_status_morgans(authed):
    col = _collection(authed, "morgan-dollars")
    a = authed.post("/api/coins", json=coin_payload(col["id"], year=1881, mint_mark="S", grade=65)).json()
    b = authed.post("/api/coins", json=coin_payload(col["id"], year=1881, mint_mark="S", grade=64, strike_designation="DMPL")).json()
    c = authed.post("/api/coins", json=coin_payload(col["id"], year=1889, mint_mark="P", grade=40, grade_prefix="XF")).json()
    sold_coin = authed.post("/api/coins", json=coin_payload(col["id"], year=1888, mint_mark="O", sale_price="50")).json()

    s = authed.get(f"/api/collections/{col['id']}/status").json()
    assert [r["label"] for r in s["rows"]] == [
        "1878 8TF", "1878 7TF, Reverse of 1878", "1878 7/8TF, Strong", "1878 7/8TF, Weak",  # 1878-P by variety
        "1880-CC", "1881-S", "1882-O", "1887-O", "1888-O", "1900-O", "1921",
    ]  # overdates (1880/79-CC, 1887/6-O) and over-mintmarks (1882-O/S, 1900-O/CC) don't get lines of their own
    assert (s["total"], s["filled"]) == (11, 1)
    row = next(r for r in s["rows"] if r["label"] == "1881-S")
    assert [x["inventory_number"] for x in row["coins"]] == [a["inventory_number"], b["inventory_number"]]  # best grade first
    assert [x["inventory_number"] for x in s["others"]] == [c["inventory_number"]]  # 1889-P isn't in the test catalog
    assert all(sold_coin["inventory_number"] not in [x["inventory_number"] for x in r["coins"]] for r in s["rows"])


def test_status_proof_vs_silver_proof(authed):
    pr = authed.get(f"/api/collections/{_collection(authed, 'kennedy-half-dollars', 'PR')['id']}/status").json()
    spr = authed.get(f"/api/collections/{_collection(authed, 'kennedy-half-dollars', 'SPR')['id']}/status").json()
    assert [r["label"] for r in pr["rows"]] == ["1964", "1992-S"]  # 1964 is silver-only, so it's a regular proof
    assert [r["label"] for r in spr["rows"]] == ["1992-S"]


def test_status_commemoratives_by_issue(authed):
    col = _collection(authed, "classic-commemoratives")
    authed.post("/api/coins", json=coin_payload(col["id"], year=1936, mint_mark="P", variety="Boone"))
    s = authed.get(f"/api/collections/{col['id']}/status").json()
    assert [r["label"] for r in s["rows"]] == [
        "1922 Grant Memorial, with Star", "1922 Grant Memorial, without Star",  # separate coins in the full set
        "1935 Boone (1935/34)", "1936 Arkansas", "1936 Boone", "1936-D Boone",
    ]
    assert [len(r["coins"]) for r in s["rows"]] == [0, 0, 0, 0, 1, 0]

    over = authed.post("/api/coins", json=coin_payload(col["id"], year=1935, mint_mark="P", variety="1935/34 Boone")).json()
    star = authed.post("/api/coins", json=coin_payload(col["id"], year=1922, mint_mark="P", variety="Grant Memorial with Star")).json()
    s = authed.get(f"/api/collections/{col['id']}/status").json()
    rows = {r["label"]: [c["inventory_number"] for c in r["coins"]] for r in s["rows"]}
    assert rows["1935 Boone (1935/34)"] == [over["inventory_number"]]  # not the plain Boone line
    assert rows["1922 Grant Memorial, with Star"] == [star["inventory_number"]]
    assert rows["1922 Grant Memorial, without Star"] == []
    assert s["modes"] == ["full", "type"] and s["mode"] == "full"


def test_export(authed):
    col = _collection(authed, "morgan-dollars", name="My Morgans")
    authed.post("/api/coins", json=coin_payload(col["id"], year=1881, mint_mark="S", price_paid="100.00", cac_sticker="gold", slab_rattler=True))
    authed.post("/api/coins", json=coin_payload(col["id"], year=1878, mint_mark="P", sale_price="200", sale_date="2026-10-01"))
    r = authed.get(f"/api/collections/{col['id']}/export")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="My_Morgans_' in r.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Coins", "Checklist", "Sold"]
    coins = list(wb["Coins"].values)
    header, rows = coins[0], coins[1:]
    assert len(rows) == 1
    rec = dict(zip(header, rows[0]))
    assert (rec["Date"], rec["Grade"], rec["CAC sticker"], rec["Collector slab"], rec["Price paid"]) == ("1881-S", "MS65", "Gold", "Rattler", 100)
    checklist = {row[0]: row[1] for row in wb["Checklist"].iter_rows(min_row=2, values_only=True) if row[0]}
    assert checklist["1881-S"] == "Yes" and checklist["1878 8TF"] == "Missing"  # sold coins don't count
    assert len(list(wb["Sold"].values)) == 2


def test_reports_require_owner(client, authed):
    col = _collection(authed, "morgan-dollars")
    client.post("/api/auth/logout")
    client.post("/api/auth/signup", json={"email": "x@example.com", "first_name": "X", "last_name": "Y", "password": "long enough pw"})
    assert client.get(f"/api/collections/{col['id']}/status").status_code == 404
    assert client.get(f"/api/collections/{col['id']}/export").status_code == 404


def test_status_washington_by_design_from_1999(authed):
    col = _collection(authed, "washington-quarters")
    de = authed.post("/api/coins", json=coin_payload(col["id"], year=1999, mint_mark="P", variety="Delaware")).json()
    pa = authed.post("/api/coins", json=coin_payload(col["id"], year=1999, mint_mark="P", pcgs_number="9604")).json()  # by PCGS #
    early = authed.post("/api/coins", json=coin_payload(col["id"], year=1934, mint_mark="P")).json()
    unknown = authed.post("/api/coins", json=coin_payload(col["id"], year=1999, mint_mark="P")).json()  # which state?

    s = authed.get(f"/api/collections/{col['id']}/status").json()
    rows = {r["label"]: [c["inventory_number"] for c in r["coins"]] for r in s["rows"]}
    # Before 1999 by date (motto varieties merged); from 1999 by design, Satin Finish merged.
    assert list(rows) == ["1934", "1998", "1999 Delaware", "1999-D Delaware", "1999 Pennsylvania"]  # release order
    assert rows["1934"] == [early["inventory_number"]]
    assert rows["1999 Delaware"] == [de["inventory_number"]]
    assert rows["1999 Pennsylvania"] == [pa["inventory_number"]]
    assert [c["inventory_number"] for c in s["others"]] == [unknown["inventory_number"]]


def test_status_washington_silver_proofs(authed):
    pr = authed.get(f"/api/collections/{_collection(authed, 'washington-quarters', 'PR')['id']}/status").json()
    spr = authed.get(f"/api/collections/{_collection(authed, 'washington-quarters', 'SPR')['id']}/status").json()
    assert [r["label"] for r in pr["rows"]] == ["1999-S Delaware", "1999-S Pennsylvania"]
    assert [r["label"] for r in spr["rows"]] == ["1999-S Delaware", "1999-S Pennsylvania"]


def test_status_1878_morgan_varieties(authed):
    col = _collection(authed, "morgan-dollars")
    by_pcgs = authed.post("/api/coins", json=coin_payload(col["id"], year=1878, mint_mark="P", pcgs_number="7074")).json()
    strong = authed.post("/api/coins", json=coin_payload(col["id"], year=1878, mint_mark="P", variety="7/8TF Strong")).json()
    eight = authed.post("/api/coins", json=coin_payload(col["id"], year=1878, mint_mark="P", variety="8tf")).json()
    unknown = authed.post("/api/coins", json=coin_payload(col["id"], year=1878, mint_mark="P")).json()

    s = authed.get(f"/api/collections/{col['id']}/status").json()
    rows = {r["label"]: [c["inventory_number"] for c in r["coins"]] for r in s["rows"] if r["year"] == 1878}
    assert rows == {
        "1878 8TF": [eight["inventory_number"]],
        "1878 7TF, Reverse of 1878": [by_pcgs["inventory_number"]],
        "1878 7/8TF, Strong": [strong["inventory_number"]],  # not mistaken for 8TF
        "1878 7/8TF, Weak": [],
    }
    assert [c["inventory_number"] for c in s["others"]] == [unknown["inventory_number"]]


def test_status_commemoratives_type_set(authed):
    col = _collection(authed, "classic-commemoratives")
    by_variety = authed.post("/api/coins", json=coin_payload(col["id"], year=1936, mint_mark="D", variety="Boone")).json()
    by_pcgs = authed.post("/api/coins", json=coin_payload(col["id"], year=1936, mint_mark="P", pcgs_number="9300")).json()
    s = authed.get(f"/api/collections/{col['id']}/status", params={"mode": "type"}).json()
    assert s["mode"] == "type"
    rows = {r["label"]: [c["inventory_number"] for c in r["coins"]] for r in s["rows"]}
    # One line per design, any date or mint; the 1935/34 overdate counts as a Boone.
    assert rows == {
        "Grant Memorial (1922)": [],  # with/without Star are one design
        "Boone (1935–1936)": [by_variety["inventory_number"]],
        "Arkansas (1936)": [by_pcgs["inventory_number"]],
    }
    assert list(rows)[0].startswith("Grant")  # earliest design first
    assert (s["total"], s["filled"]) == (3, 2)

    # Series without a type set fall back to the full set
    m = _collection(authed, "morgan-dollars")
    assert authed.get(f"/api/collections/{m['id']}/status", params={"mode": "type"}).json()["mode"] == "full"

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(authed.get(f"/api/collections/{col['id']}/export").content))
    assert "Type Set" in wb.sheetnames and wb["Type Set"]["A1"].value == "Design"


def test_gold_type_set(authed, monkeypatch):
    import test_catalog as fake
    from app.series import BY_SLUG

    # One fake catalog coin per gold design section, plus two $20 Liberty Type 2 dates.
    nodes = dict(fake.NODES)
    for i, (label, node) in enumerate(BY_SLUG["gold-type-set"].type_rows):
        year = int(label.split("(")[1][:4])
        nodes[node] = [fake.item(1000 + i, str(70000 + i), str(year), denom="$")]
    nodes[10340] = [fake.item(2000, "8950", "1866-S", denom="$20"), fake.item(2001, "8960", "1870-CC", denom="$20")]
    # Each member series has a proof section too; give them something so the sync succeeds.
    for s in {BY_SLUG[k] for k in BY_SLUG if BY_SLUG[k].category == "Gold"}:
        for n in s.pr_nodes:
            nodes.setdefault(n, [fake.item(5000 + n, None, "1900", strike="PR", denom="$")])
    monkeypatch.setattr(fake, "NODES", nodes)

    types = {t["slug"]: t for t in authed.get("/api/coin-types").json()}
    col = authed.post("/api/collections", json={"name": "My Gold Type Set", "coin_type_id": types["gold-type-set"]["id"]}).json()
    assert col["collection_type_slug"] == "gold-type-set" and col["collection_type_label"] == "Gold Type Set"

    # The coin's series must be chosen, and must be one of the type set's members
    assert authed.post("/api/coins", json=coin_payload(col["id"], year=1870, mint_mark="CC")).status_code == 422
    morgan = coin_payload(col["id"], year=1881, mint_mark="S", coin_type_id=types["morgan-dollars"]["id"])
    assert authed.post("/api/coins", json=morgan).status_code == 422

    t2 = authed.post("/api/coins", json=coin_payload(col["id"], year=1870, mint_mark="CC", coin_type_id=types["liberty-double-eagles"]["id"])).json()
    assert t2["inventory_number"] == "G20L-00001" and t2["coin_type_id"] == types["liberty-double-eagles"]["id"]
    saint = authed.post("/api/coins", json=coin_payload(col["id"], year=1907, mint_mark="P", pcgs_number=str(70000 + 22), coin_type_id=types["saint-gaudens-double-eagles"]["id"])).json()

    s = authed.get(f"/api/collections/{col['id']}/status").json()
    assert s["total"] == 23 and s["filled"] == 2
    rows = {r["label"]: [c["inventory_number"] for c in r["coins"]] for r in s["rows"]}
    assert rows["$20 Liberty Type 2 (1866–1876)"] == [t2["inventory_number"]]  # by series + year
    assert rows["$20 Saint-Gaudens (1907–1933)"] == [saint["inventory_number"]]  # by PCGS #
    assert list(rows)[0] == "$1 Gold Type 1 (1849–1854)"  # checklist order

    # A coin can't switch series once added
    body = coin_payload(col["id"], year=1870, mint_mark="CC", coin_type_id=types["liberty-eagles"]["id"])
    assert authed.put(f"/api/coins/{t2['inventory_number']}", json=body).status_code == 400


def test_insurance_pdf(authed):
    from PIL import Image as PILImage

    col = _collection(authed, "morgan-dollars", name="My Morgans")
    coin = authed.post("/api/coins", json=coin_payload(col["id"], year=1881, mint_mark="S", price_paid="100.00", cac_sticker="green")).json()
    authed.post("/api/coins", json=coin_payload(col["id"], year=1878, mint_mark="P", sale_price="200", sale_date="2026-10-01"))
    buf = io.BytesIO()
    PILImage.new("RGB", (600, 400), (180, 180, 190)).save(buf, "PNG")
    files = [("files", ("good.png", buf.getvalue(), "image/png")), ("files", ("bad.png", b"not an image", "image/png"))]
    authed.post(f"/api/coins/{coin['inventory_number']}/photos", files=files)

    for url in ("/api/reports/insurance", f"/api/reports/insurance?collection_id={col['id']}"):
        r = authed.get(url)
        assert r.status_code == 200, r.text
        assert r.headers["content-type"] == "application/pdf"
        assert r.content.startswith(b"%PDF")
        assert "Insurance_Inventory_" in r.headers["content-disposition"]
    assert authed.get("/api/reports/insurance?collection_id=99999").status_code == 404


def test_export_chosen_columns(authed):
    col = _collection(authed, "morgan-dollars", name="My Morgans")
    authed.post("/api/coins", json=coin_payload(col["id"], year=1881, mint_mark="S", price_paid="100.00"))
    authed.post("/api/coins", json=coin_payload(col["id"], year=1878, mint_mark="P", sale_price="200", sale_date="2026-10-01"))
    cols = authed.get("/api/reports/export-columns").json()
    names = [c["name"] for c in cols]
    assert names[0] == "Inventory #" and "Sale price" in names
    assert next(c for c in cols if c["name"] == "Sale price")["sold_only"]

    r = authed.get(f"/api/collections/{col['id']}/export", params=[("columns", "Date"), ("columns", "Price paid"), ("columns", "Sale price")])
    wb = load_workbook(io.BytesIO(r.content))
    assert list(wb["Coins"].values)[0] == ("Date", "Price paid")
    assert list(wb["Sold"].values)[0] == ("Date", "Price paid", "Sale price")
    assert authed.get(f"/api/collections/{col['id']}/export", params={"columns": "Nope"}).status_code == 400


def test_status_overdates_fill_the_date(authed):
    col = _collection(authed, "morgan-dollars")
    od = authed.post("/api/coins", json=coin_payload(col["id"], year=1880, mint_mark="CC", variety="1880/79")).json()
    om = authed.post("/api/coins", json=coin_payload(col["id"], year=1900, mint_mark="O", variety="O/CC")).json()
    s = authed.get(f"/api/collections/{col['id']}/status").json()
    rows = {r["label"]: [c["inventory_number"] for c in r["coins"]] for r in s["rows"]}
    assert rows["1880-CC"] == [od["inventory_number"]] and rows["1900-O"] == [om["inventory_number"]]
    assert not any("/" in label and not label.startswith("1878") for label in rows)
    assert s["others"] == []


def test_upgrade_finder(authed):
    col = _collection(authed, "morgan-dollars", name="My Morgans")
    coin = authed.post("/api/coins", json=coin_payload(col["id"], year=1881, mint_mark="S", grade=64, pcgs_number="7130")).json()
    authed.post("/api/coins", json=coin_payload(col["id"], year=1888, mint_mark="O", grade=63))  # never priced
    d = authed.get(f"/api/collections/{col['id']}/upgrades").json()
    assert d["unpriced"] == ["1881-S", "1888-O"] and d["items"] == []  # nothing priced yet
    authed.post(f"/api/coins/{coin['inventory_number']}/cpg")
    d = authed.get(f"/api/collections/{col['id']}/upgrades").json()
    assert d["unpriced"] == ["1888-O"]
    [item] = d["items"]
    assert (item["label"], item["current_value"], item["next_label"], item["next_cpg"], item["cost"], item["pct"]) == (
        "1881-S", "900.00", "65", "1150.00", "250.00", 28)


def test_upgrade_finder_plus_grade(authed, monkeypatch):
    from types import SimpleNamespace

    from app.db import SessionLocal
    from app.models import Coin, GradePriceSheet
    from app.services import upgrades

    col = _collection(authed, "morgan-dollars", name="Plus")
    authed.post("/api/coins", json=coin_payload(col["id"], year=1881, mint_mark="S", grade=65, grade_plus=True, pcgs_number="7130"))
    with SessionLocal() as db:
        db.add(GradePriceSheet(key="pcgs:7130", name="1881-S $1", rows=[
            {"grade": 65, "label": "MS65", "cac": False, "cpg": "500"},
            {"grade": 65, "label": "MS65+", "cac": False, "cpg": "700"},
            {"grade": 66, "label": "MS66", "cac": False, "cpg": "900"}]))
        db.commit()

    def one_row(db, col, mode):  # the checklist reduced to this coin's line
        row = SimpleNamespace(coins=[db.query(Coin).one()], pcgs_numbers={"7130"}, label="1881-S", year=1881,
                              mint_mark="S", variety=None)
        return SimpleNamespace(rows=[row])

    monkeypatch.setattr(upgrades.checklist, "build", one_row)
    [item] = authed.get(f"/api/collections/{col['id']}/upgrades").json()["items"]
    assert (item["current_value"], item["next_label"], item["cost"]) == ("700.00", "MS66", "200.00")  # 65+ priced as 65+
