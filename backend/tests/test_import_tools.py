"""Import template, the optional Stored in / Status columns, and the spreadsheet cleanup tool."""

import io
import json

from openpyxl import load_workbook
from test_api import coin_payload


def _csv(text: str):
    return {"file": ("coins.csv", text.encode(), "text/csv")}


def test_template_imports_cleanly(authed):
    r = authed.get("/api/import/template")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Coins", "Instructions"]
    p = authed.post("/api/import/preview", files={"file": ("t.xlsx", r.content, "application/octet-stream")}).json()
    assert p["error_count"] == 0, [x["errors"] for x in p["rows"]]
    out = authed.post("/api/import/commit", json={"collection_id": None, "rows": p["rows"]}).json()
    assert len(out["created"]) == 4
    guide = authed.get("/api/import/columns").json()
    assert next(c for c in guide if c["header"] == "Date")["need"] == "Required"


def test_storage_and_status_columns(authed, morgan_collection):
    sub = authed.post("/api/grading/submissions", json={"company": "NGC", "name": "Morgans out"}).json()
    authed.post("/api/storage/spots", json={"name": "Box #12"})
    text = ("Type,Date,Grading service,Cert #,Grade,Stored in,Position,Status,Submission,Submitted to\n"
            "Morgan Dollar,1881-S,PCGS,11111111,MS64,box #12,Row 1,,,\n"
            "Morgan Dollar,1882-S,PCGS,22222222,MS64,Album B,,At Grading,Morgans out,\n"
            "Morgan Dollar,1883-O,PCGS,33333333,MS63,,,At Grading,,PCGS\n"
            "Morgan Dollar,1884-O,PCGS,44444444,MS63,,,At Grading,#99,\n"
            "Morgan Dollar,1885-O,PCGS,55555555,MS63,,,Lost,,\n")
    p = authed.post("/api/import/preview", files=_csv(text)).json()
    rows = {r["row"]: r for r in p["rows"]}
    assert rows[2]["placement"] == "Stored in box #12 · Row 1" and not rows[2]["errors"]
    assert "Album B (new)" in rows[3]["placement"] and "At grading: #1 NGC" in rows[3]["placement"]
    assert "new PCGS submission" in rows[4]["placement"]
    assert "No grading submission" in rows[5]["errors"][0]
    assert "Status should be" in rows[6]["errors"][0]
    out = authed.post("/api/import/commit", json={"collection_id": None, "rows": p["rows"]}).json()
    c1, c2, c3 = (authed.get(f"/api/coins/{inv}").json() for inv in out["created"])
    assert (c1["storage_name"], c1["storage_position"], c1["location"]) == ("Box #12", "Row 1", "inventory")
    assert (c2["storage_name"], c2["location"], c2["grading"]["id"]) == ("Album B", "grading", sub["id"])
    assert c3["grading"]["company"] == "PCGS" and c3["storage_id"] is None
    # No such columns: nothing is set.
    p = authed.post("/api/import/preview", files=_csv("Type,Date,Grade\nMorgan Dollar,1886,MS62\n")).json()
    assert p["rows"][0]["placement"] is None


def test_cleanup(authed, morgan_collection):
    authed.post("/api/coins", json=coin_payload(morgan_collection["id"], cert_number="99999999"))
    text = ("Coin,Grade,Paid,Bought on,Where,My comments\n"
            '1881-S Morgan Dollar PCGS MS65+ DMPL CAC Green #12345678,,"$1,250",3/15/25,Heritage,nice\n'
            "1921 peace dollar ngc 1234567-001,ms 64,250,2025-04-01,,\n"
            "1880 Mystery Coin,AU50,10,,,\n"
            "1882 Morgan PCGS MS63 99999999,,,,,\n")
    mapping = json.dumps({"Where": "purchased_from", "My comments": "notes_append"})
    r = authed.post("/api/import/cleanup", files=_csv(text), data={"mapping": mapping}).json()
    heads = {h["header"]: h["field"] for h in r["headers"]}
    assert heads["Coin"] == "date" and heads["Paid"] == "price_paid" and heads["Where"] == "purchased_from"
    a, b, c, d = r["rows"]
    assert {k: a["cleaned"][k] for k in ("Type", "Date", "Grading service", "Cert #", "Grade", "CAC sticker",
                                         "Price paid", "Date purchased", "Purchased from", "Notes")} == {
        "Type": "Morgan Dollar", "Date": "1881-S", "Grading service": "PCGS", "Cert #": "12345678",
        "Grade": "MS65+ DMPL", "CAC sticker": "Green", "Price paid": "1250.00", "Date purchased": "2025-03-15",
        "Purchased from": "Heritage", "Notes": "My comments: nice"}
    assert any(ch["column"] == "Date" and ch["before"].startswith("1881-S Morgan") for ch in a["changes"])
    assert (b["cleaned"]["Type"], b["cleaned"]["Grade"], b["cleaned"]["Grading service"]) == ("Peace Dollar", "MS64", "NGC")
    assert c["problems"] and not a["problems"]
    assert any("Already in your collection" in p for p in d["problems"])
    assert r["counts"] == {"rows": 4, "changed": 4, "problems": 2}

    f = authed.post("/api/import/cleanup/file", files=_csv(text), data={"mapping": mapping})
    assert f.status_code == 200 and "coins_cleaned.xlsx" in f.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(f.content))
    assert wb.sheetnames == ["Coins", "Changes", "Problems"]
    # The cleaned file imports: the two good rows are ready.
    p = authed.post("/api/import/preview", files={"file": ("c.xlsx", f.content, "application/octet-stream")}).json()
    assert p["ok_count"] == 2
    assert authed.post("/api/import/cleanup", files=_csv(text), data={"mapping": '{"Coin": "rm -rf"}'}).status_code == 422


def test_cleanup_grade_cells_with_service_cert_and_designation(authed):
    """What people write in a Grade column: "FR02 (32611221 PCGS)", "MS64PL", "PF 69 Ultra Cameo NGC 2345678-004";
    a cert with no service says which one it is by its format."""
    text = ("Coin,Grade,Cert\n"
            '1878-P Morgan Dollar (7TF),FR02 (32611221 PCGS),\n'
            "1881-S Morgan Dollar,MS64PL,\n"
            "1895 Morgan Dollar,PF 69 Ultra Cameo NGC 2345678-004,\n"
            "1884-CC Morgan Dollar,MS63 Toned CAC Green,23456789\n"
            "1921 Morgan Dollar,AU58,1234567-002\n"
            "1921 Morgan Dollar,AU58,9876543210\n")
    rows = authed.post("/api/import/cleanup", files=_csv(text)).json()["rows"]
    got = [(r["cleaned"]["Grading service"], r["cleaned"]["Cert #"], r["cleaned"]["Grade"]) for r in rows]
    assert got[:5] == [("PCGS", "32611221", "FR2"), ("Raw", "", "MS64 PL"), ("NGC", "2345678-004", "PR69 DCAM"),
                       ("PCGS", "23456789", "MS63"), ("NGC", "1234567-002", "AU58")]
    assert rows[0]["cleaned"]["Variety"] == "7TF" and not rows[0]["problems"]
    assert rows[3]["cleaned"]["CAC sticker"] == "Green" and "Grade note: Toned" in rows[3]["cleaned"]["Notes"]
    assert rows[3]["notes"] and "cert number" in rows[3]["notes"][0]  # PCGS from an 8-digit cert: a guess, flagged
    assert not rows[4]["notes"]  # NGC's format is certain
    assert any("no grading service" in p for p in rows[5]["problems"])  # 10 digits: can't tell
